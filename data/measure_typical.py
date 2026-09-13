"""Measure, on the real data, whether "the average is not typical" holds -- and
in what form -- BEFORE any of it is written into a page.

Two candidate lessons, and the data decides which one each dataset supports:

  A. SKEWED: most observations fall BELOW the mean, so the average describes
     almost nobody and overstates the typical case.
  B. BIMODAL: the mean lands in the VALLEY between two clusters, so the
     average is a value that rarely occurs at all.

The filters are re-derived here from the raw source and cross-checked against
the population means already recorded in datasets.js. If my filtering drifts
from what the pages ship, that cross-check fails loudly rather than producing a
figure that quietly describes a different population.
"""
from __future__ import annotations

import csv
import io
import json
import statistics
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
CACHE = HERE / "cache"
ROOT = HERE.parent


def shipped_stats() -> dict[str, dict]:
    s = (ROOT / "datasets.js").read_text(encoding="utf-8")
    arr = json.loads(s[s.index("["):s.rindex("]") + 1])
    return {d["id"]: d for d in arr}


# ------------------------------------------------------------------ loaders
def load_retail() -> list[float]:
    """Invoice totals, with the same filters build_datasets.py applies."""
    from openpyxl import load_workbook
    raw = (CACHE / "online_retail.zip").read_bytes()
    zf = zipfile.ZipFile(io.BytesIO(raw))
    inner = [n for n in zf.namelist() if n.lower().endswith((".xlsx", ".xls"))][0]
    wb = load_workbook(io.BytesIO(zf.read(inner)), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    col = {h: i for i, h in enumerate(header)}
    NONPRODUCT = {"POST", "DOT", "C2", "M", "BANK CHARGES", "PADS", "AMAZONFEE",
                  "S", "D", "CRUK", "B", "GIFT"}
    totals: dict[str, float] = {}
    for r in rows:
        inv = r[col["InvoiceNo"]]
        if inv is None:
            continue
        inv = str(inv).strip()
        if inv.upper().startswith("C"):
            continue
        if str(r[col["StockCode"]] or "").strip().upper() in NONPRODUCT:
            continue
        try:
            q = float(r[col["Quantity"]]); p = float(r[col["UnitPrice"]])
        except (TypeError, ValueError):
            continue
        if q <= 0 or p <= 0:
            continue
        totals[inv] = totals.get(inv, 0.0) + q * p
    wb.close()
    return sorted(totals.values())


def load_csv_col(name: str, column: str) -> list[float]:
    with (CACHE / name).open(newline="", encoding="utf-8") as f:
        return sorted(float(row[column]) for row in csv.DictReader(f)
                      if row.get(column) not in (None, ""))


# ------------------------------------------------------------------ measures
def below_mean(vals: list[float], mu: float) -> int:
    """Count strictly below the mean. vals must be sorted."""
    lo, hi = 0, len(vals)
    while lo < hi:
        mid = (lo + hi) // 2
        if vals[mid] < mu:
            lo = mid + 1
        else:
            hi = mid
    return lo


def histogram(vals: list[float], nbins: int) -> tuple[list[int], float, float]:
    lo, hi = vals[0], vals[-1]
    w = (hi - lo) / nbins
    counts = [0] * nbins
    for v in vals:
        k = int((v - lo) / w)
        counts[min(k, nbins - 1)] += 1
    return counts, lo, w


def peaks(counts: list[int], min_gap: int) -> list[int]:
    """Well-SEPARATED local maxima of a smoothed histogram.

    A naive local-maximum search on the raw counts is wrong here, and produced a
    false claim on the first run: it returned bins at 77.2 and 81.6 minutes as
    "the two clusters", when both of those sit inside the SAME right-hand
    cluster -- that mode is broad and noisy and has several local maxima.
    Smoothing over one bin either side removes the noise, and requiring a
    minimum gap stops two bumps in one hump being reported as two clusters.
    """
    sm = []
    for i in range(len(counts)):
        a = max(0, i - 1)
        b = min(len(counts), i + 2)
        sm.append(sum(counts[a:b]) / (b - a))
    cand = [i for i in range(1, len(sm) - 1)
            if sm[i] > sm[i - 1] and sm[i] >= sm[i + 1]]
    cand.sort(key=lambda i: -sm[i])
    chosen: list[int] = []
    for i in cand:
        if all(abs(i - j) >= min_gap for j in chosen):
            chosen.append(i)
    return chosen


def within(vals: list[float], centre: float, half: float) -> int:
    return sum(1 for v in vals if abs(v - centre) <= half)


# ------------------------------------------------------------------ main
def main() -> None:
    ship = shipped_stats()
    fails: list[str] = []

    print("loading the real sources (this reads the 23 MB retail workbook)...")
    data = {
        "retail":  (load_retail(), "retail", "£", "", "order value per invoice"),
        "diamond": (load_csv_col("diamonds.csv", "price"), "diamond", "$", "",
                    "price per diamond"),
        "geyser":  (load_csv_col("geyser.csv", "waiting"), "geyser", "", " min",
                    "wait between eruptions"),
    }

    print("\ncross-check: my filters must reproduce the population mean the")
    print("pages already ship, otherwise these figures describe a different set")
    for key, (vals, sid, _, _, _) in data.items():
        mu = statistics.fmean(vals)
        want = ship[sid]["full_mu"]
        ok = abs(mu - want) < 0.01
        print(f"  {key:<8} n={len(vals):>6,}  mean {mu:>10.4f}  "
              f"datasets.js {want:>10.4f}  {'MATCH' if ok else 'MISMATCH'}")
        if not ok:
            fails.append(f"{key}: my mean {mu:.4f} != shipped {want:.4f} — filters differ")
        if len(vals) != ship[sid]["n_source"]:
            fails.append(f"{key}: n {len(vals)} != shipped n_source {ship[sid]['n_source']}")

    print("\n" + "=" * 74)
    print("LESSON A — is the average above what most people actually do?")
    print("=" * 74)
    results = {}
    for key, (vals, sid, pre, suf, what) in data.items():
        n = len(vals)
        mu = statistics.fmean(vals)
        med = statistics.median(vals)
        nb = below_mean(vals, mu)
        pct = 100.0 * nb / n
        ratio = mu / med if med else float("nan")
        print(f"\n  {key} — {what}")
        print(f"    mean            {pre}{mu:,.2f}{suf}")
        print(f"    median          {pre}{med:,.2f}{suf}   (mean is {ratio:.2f}x the median)")
        print(f"    below the mean  {nb:,} of {n:,} = {pct:.1f}%")
        print(f"    the mean sits at the {pct:.1f}th percentile")
        results[key] = dict(n=n, mu=mu, median=med, below=nb, pct=pct, ratio=ratio)

    print("\n" + "=" * 74)
    print("LESSON B — for a two-cluster quantity, is the average a rare value?")
    print("=" * 74)
    gv = data["geyser"][0]
    mu_g = statistics.fmean(gv)
    counts, lo, w = histogram(gv, 24)
    pk = peaks(counts, min_gap=7)   # >= ~15 min apart, so one broad
                                   # hump cannot supply both modes
    print(f"\n  histogram of {len(gv)} waits, {len(counts)} bins of {w:.2f} min")
    for i, c in enumerate(counts):
        centre = lo + (i + 0.5) * w
        mark = "  <- mean" if abs(centre - mu_g) <= w / 2 else ""
        star = " *PEAK" if i in pk[:2] else ""
        print(f"    {lo + i * w:5.1f}–{lo + (i + 1) * w:5.1f}  "
              f"{'#' * max(0, round(c / 2)):<22} {c:>3}{star}{mark}")
    if len(pk) < 2:
        fails.append("geyser: expected two clusters, found fewer — lesson B does not hold")
    else:
        p1, p2 = sorted(pk[:2])
        c1 = lo + (p1 + 0.5) * w
        c2 = lo + (p2 + 0.5) * w
        half = w
        n_mu, n_c1, n_c2 = (within(gv, mu_g, half), within(gv, c1, half),
                            within(gv, c2, half))
        print(f"\n  the two clusters sit at about {c1:.1f} min and {c2:.1f} min")
        print(f"  the mean wait is {mu_g:.1f} min — between them")
        print(f"  waits within +/-{half:.1f} min of ...")
        print(f"    the mean      ({mu_g:5.1f}): {n_mu:>3} of {len(gv)} = "
              f"{100 * n_mu / len(gv):.1f}%")
        print(f"    cluster one   ({c1:5.1f}): {n_c1:>3} of {len(gv)} = "
              f"{100 * n_c1 / len(gv):.1f}%")
        print(f"    cluster two   ({c2:5.1f}): {n_c2:>3} of {len(gv)} = "
              f"{100 * n_c2 / len(gv):.1f}%")
        rarer = n_mu < n_c1 and n_mu < n_c2
        print(f"\n  is the average rarer than either cluster? "
              f"{'YES' if rarer else 'NO'}")
        results["geyser_bimodal"] = dict(mu=mu_g, c1=c1, c2=c2, half=half,
                                         n_mu=n_mu, n_c1=n_c1, n_c2=n_c2,
                                         rarer=rarer, n=len(gv))
        if not rarer:
            fails.append("geyser: the mean is NOT rarer than the clusters — "
                         "lesson B must be reworded or dropped")

    print("\n" + "=" * 74)
    print("VERDICT")
    print("=" * 74)
    for key in ("retail", "diamond", "geyser"):
        r = results[key]
        holds = r["pct"] > 55
        print(f"  {key:<8} {r['pct']:.1f}% below average — "
              f"{'supports lesson A' if holds else 'does NOT support lesson A'}")
    gb = results.get("geyser_bimodal")
    if gb:
        print(f"  geyser   average occurs {100 * gb['n_mu'] / gb['n']:.1f}% of the time vs "
              f"{100 * max(gb['n_c1'], gb['n_c2']) / gb['n']:.1f}% at the busier cluster — "
              f"{'supports lesson B' if gb['rarer'] else 'does NOT support lesson B'}")

    (HERE / "measured_typical.json").write_text(json.dumps(results, indent=1),
                                                encoding="utf-8")
    print(f"\nwritten to {HERE / 'measured_typical.json'}")

    if fails:
        print("\nPROBLEMS:")
        for f in fails:
            print("  -", f)
        sys.exit(1)


if __name__ == "__main__":
    main()
